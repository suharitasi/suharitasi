#!/usr/bin/env python3
"""DSİ 2024 resmi istatistik Excel'lerinden havza veri künyesini üretir.

Girdi:  kaynak/dsi/1.2.*.xlsx (yüzeysuyu), kaynak/dsi/1.3.*.xlsx (YAS)
Çıktı:  data/havza-veri.json

Her alan kaynağıyla işaretlenir; bulunamayan alan null bırakılır ve
arayüzde "veri yok (tarih)" gösterilir. UYDURMA DEĞER YAZILMAZ.
"""
import json
import unicodedata

import openpyxl

YUZEY = 'kaynak/dsi/1.2.havzalara_gore_yillik_ortalama_yuzeysuyu_su_potansiyeli_20132024.xlsx'
YAS = 'kaynak/dsi/1.3.havzalara_gore_yillik_yeraltisuyu_potansiyeli_20132024.xlsx'
CIKTI = 'data/havza-veri.json'

# SYGM havza koruma eylem planı PDF'leri — 2026-07-14'te HTTP 200 doğrulandı.
HKEP_TABAN = 'https://www.tarimorman.gov.tr/SYGM/Belgeler/havza%20koruma%20eylem%20planlar%C4%B1/'
EYLEM_PLANLARI = {
    '01': None,  # HKEP bulunamadı; NHYP mevcut (aşağıda ayrıca işlenir)
    '02': HKEP_TABAN + 'Marmara_Havzasi.pdf',
    '03': HKEP_TABAN + 'Susurluk-Havzasi.pdf',
    '04': HKEP_TABAN + 'Kuzey_Ege_Havzasi.pdf',
    '05': HKEP_TABAN + 'Gediz_web.pdf',
    '06': HKEP_TABAN + 'Kucuk_Menderes_Havzasi.pdf',
    '07': HKEP_TABAN + 'B.Menderes_Havzas%C4%B1.pdf',
    '08': HKEP_TABAN + 'Bati_Akdeniz_web.pdf',
    '09': HKEP_TABAN + 'Antalya_web.pdf',
    '10': HKEP_TABAN + 'Burdur_Havzas%C4%B1.pdf',
    '11': HKEP_TABAN + 'Akarcay_web.pdf',
    '12': HKEP_TABAN + 'Sakarya_web.pdf',
    '13': HKEP_TABAN + 'Bat%C4%B1%20Karadeniz_web_rev3.pdf',
    '14': HKEP_TABAN + 'Ye%C5%9Fil%C4%B1rmak%20Havzas%C4%B1.pdf',
    '15': HKEP_TABAN + 'K%C4%B1z%C4%B1l%C4%B1rmak_Havzas%C4%B1.pdf',
    '16': HKEP_TABAN + 'Konya_Kapali_Havzasi.pdf',
    '17': HKEP_TABAN + 'Dogu_Akdeniz_web.pdf',
    '18': HKEP_TABAN + 'Seyhan_Havzasi.pdf',
    '19': None,  # bulunamadı (2026-07-14)
    '20': HKEP_TABAN + 'Ceyhan_Havzas%C4%B1.pdf',
    '21': None,  # Fırat-Dicle — bulunamadı (2026-07-14)
    '22': HKEP_TABAN + 'Dogu_Karadeniz_web.pdf',
    '23': None,  # Çoruh — bulunamadı (2026-07-14)
    '24': None,  # Aras — bulunamadı (2026-07-14)
    '25': HKEP_TABAN + 'Van_Golu_web.pdf',
}
NHYP = {
    '01': 'https://www.tarimorman.gov.tr/SYGM/Belgeler/NHYP%20DEN%C4%B0Z/'
          'MER%C4%B0%C3%87-ERGENE%20NEH%C4%B0R%20HAVZASI%20Y%C3%96NET%C4%B0M%20PLANI.pdf',
}

AD_STANDART = {
    'Meriç Ergene': 'Meriç-Ergene Havzası',
    'Meriç - Ergene': 'Meriç-Ergene Havzası',
    'Marmara': 'Marmara Havzası',
    'Susurluk': 'Susurluk Havzası',
    'Kuzey Ege': 'Kuzey Ege Havzası',
    'Gediz': 'Gediz Havzası',
    'Küçük Menderes': 'Küçük Menderes Havzası',
    'Büyük Menderes': 'Büyük Menderes Havzası',
    'Batı Akdeniz': 'Batı Akdeniz Havzası',
    'Antalya': 'Antalya Havzası',
    'Burdur Göller': 'Burdur Havzası',
    'Akarçay': 'Akarçay Havzası',
    'Sakarya': 'Sakarya Havzası',
    'Batı Karadeniz': 'Batı Karadeniz Havzası',
    'Yeşilırmak': 'Yeşilırmak Havzası',
    'Kızılırmak': 'Kızılırmak Havzası',
    'Konya Kapalı': 'Konya Kapalı Havzası',
    'Doğu Akdeniz': 'Doğu Akdeniz Havzası',
    'Seyhan': 'Seyhan Havzası',
    'Asi': 'Asi Havzası',
    'Ceyhan': 'Ceyhan Havzası',
    'Fırat - Dicle (*)(**)': 'Fırat-Dicle Havzası',
    'Fırat - Dicle': 'Fırat-Dicle Havzası',
    'Doğu Karadeniz': 'Doğu Karadeniz Havzası',
    'Çoruh': 'Çoruh Havzası',
    'Aras': 'Aras Havzası',
    'Van Gölü': 'Van Gölü Havzası',
}


def adla(s):
    return AD_STANDART.get(unicodedata.normalize('NFC', str(s).strip()), None)


def son_yil_blok(ws, blok_genislik):
    """Başlık satırındaki en büyük yılın sütun aralığını döndürür."""
    yillar = {}
    for satir in ws.iter_rows(min_row=1, max_row=6):
        for h in satir:
            if isinstance(h.value, (int, float)) and 2000 < h.value < 2100:
                yillar[int(h.value)] = h.column
            elif isinstance(h.value, str) and h.value.strip().isdigit() and 2000 < int(h.value) < 2100:
                yillar[int(h.value)] = h.column
    yil = max(yillar)
    return yil, yillar[yil]


havzalar = {}

# 1.2 — yüzeysuyu: yıl bloğu 3 sütun (yağış alanı, akış, iştirak %)
wb = openpyxl.load_workbook(YUZEY, data_only=True)
ws = wb.active
yil12, kolon = son_yil_blok(ws, 3)
for r in ws.iter_rows(min_row=5, values_only=False):
    ad = adla(r[2].value) if len(r) > 2 else None
    if not ad:
        continue
    no = str(r[1].value).zfill(2)
    alan = ws.cell(r[0].row, kolon).value
    akis = ws.cell(r[0].row, kolon + 1).value
    havzalar[no] = {
        'no': no,
        'ad': ad,
        'yagisAlani_km2': alan,
        'yuzeysuyuPotansiyeli_km3': akis,
        'yuzeysuyuYili': yil12,
    }

# 1.3 — YAS: yıl bloğu 2 sütun (beslenim, işletme rezervi)
wb = openpyxl.load_workbook(YAS, data_only=True)
ws = wb.active
yil13, kolon = son_yil_blok(ws, 2)
for r in ws.iter_rows(min_row=5, values_only=False):
    ad = adla(r[2].value) if len(r) > 2 else None
    if not ad:
        continue
    no = str(r[1].value).zfill(2)
    beslenim = ws.cell(r[0].row, kolon).value
    rezerv = ws.cell(r[0].row, kolon + 1).value
    havzalar.setdefault(no, {'no': no, 'ad': ad})
    havzalar[no]['yasBeslenimi_hm3'] = round(beslenim, 1) if isinstance(beslenim, float) else beslenim
    havzalar[no]['yasIsletmeRezervi_hm3'] = round(rezerv, 1) if isinstance(rezerv, float) else rezerv
    havzalar[no]['yasYili'] = yil13

for no, h in havzalar.items():
    h['eylemPlani'] = EYLEM_PLANLARI.get(no)
    h['nehirHavzasiYonetimPlani'] = NHYP.get(no)
    h['tahsis'] = None  # açık, havza bazlı tahsis verisi bulunamadı (2026-07-14)

sonuc = {
    'aciklama': 'Türkiye 25 su havzası veri künyesi. Boş (null) alan = veri yok; uydurma değer yazılmaz.',
    'kaynaklar': {
        'yuzeysuyu': {
            'ad': 'DSİ 2024 Yılı Resmi Su Kaynakları İstatistikleri, Tablo 1.2 (Havzalara Göre Yıllık Ortalama Yüzeysuyu Potansiyeli)',
            'url': 'https://www.dsi.gov.tr/Sayfa/Detay/2186',
            'indirmeTarihi': '2026-07-14',
        },
        'yas': {
            'ad': 'DSİ 2024 Yılı Resmi Su Kaynakları İstatistikleri, Tablo 1.3 (Havzalara Göre Yıllık Yeraltısuyu Potansiyeli)',
            'url': 'https://www.dsi.gov.tr/Sayfa/Detay/2186',
            'indirmeTarihi': '2026-07-14',
        },
        'eylemPlanlari': {
            'ad': 'SYGM Havza Koruma Eylem Planları',
            'url': 'https://www.tarimorman.gov.tr/SYGM/Sayfalar/Detay.aspx?SayfaId=6',
            'dogrulamaTarihi': '2026-07-14',
        },
    },
    'havzalar': [havzalar[k] for k in sorted(havzalar)],
}

json.dump(sonuc, open(CIKTI, 'w', encoding='utf-8'), ensure_ascii=False, indent=2)
print('havza:', len(havzalar), '| yüzeysuyu yılı:', yil12, '| YAS yılı:', yil13)
eksik = [h['ad'] for h in havzalar.values() if not h.get('eylemPlani') and not h.get('nehirHavzasiYonetimPlani')]
print('eylem planı bulunamayan:', eksik)
