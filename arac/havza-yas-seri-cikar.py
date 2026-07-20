#!/usr/bin/env python3
"""DSİ YAS xlsx arşivinden havza potansiyel bloğu için TEK normalize kaynak üretir.

Girdi (data/arsiv/dsi-yas/2026-07-20/2024-seti/):
  1.3  havza × yıl (2013-2024): beslenim (potansiyel) + işletme rezervi
  1.5  ulusal işletme rezervi serisi (1995-2024)
  1.6  ulusal sektör tahsisi (1995-2024)
Çıktı: data/canli/havza-yas.json  (25 havza sayfası bundan beslenir — çift bakım yok)

Kurallar: xlsx'te olmayan yıl/havza değeri atlanır (null yazılmaz, uydurma yok).
Havza eşleştirme resmî 'no' (1-25) üzerinden; havza-veri.json kanonik ad kaynağı.
Eşleşmeyen 'eslesmedi' listesine düşer — tahminle eşleştirme yok.
"""
import json
import openpyxl

A = "data/arsiv/dsi-yas/2026-07-20/2024-seti"
F13 = f"{A}/1.3.havzalara_gore_yillik_yeraltisuyu_potansiyeli_20132024.xlsx"
F15 = f"{A}/1.5.yeraltisuyu_isletme_rezervi_19952024.xlsx"
F16 = f"{A}/1.6.sektor_bazinda_yeraltisuyundan_yapilan_tahsis_miktari_19952024.xlsx"
CIKTI = "data/canli/havza-yas.json"

URL = "https://www.dsi.gov.tr/Sayfa/Detay/2186"
TARIH = "2026-07-20"


def r3(v):
    return round(float(v), 3) if isinstance(v, (int, float)) else None


# kanonik no -> ad (mevcut şema)
kanonik = {h["no"]: h["ad"] for h in json.load(open("data/havza-veri.json"))["havzalar"]}

# --- 1.3: yıl kolonları (row4), her yıl için beslenim=col, rezerv=col+1 ---
wb = openpyxl.load_workbook(F13, data_only=True)
ws = wb.active
yil_kol = {}
for c in range(2, 28):  # ilk blok (col29+ ikinci tekrar bloğu; alınmaz)
    v = ws.cell(4, c).value
    if isinstance(v, int) and 2013 <= v <= 2024:
        yil_kol[v] = c
yillar = sorted(yil_kol)

havzalar = {}
eslesmedi = []
for r in range(5, ws.max_row + 1):
    no_raw = ws.cell(r, 2).value
    if not isinstance(no_raw, int) or not (1 <= no_raw <= 25):
        continue
    no = f"{no_raw:02d}"
    ad_kaynak = str(ws.cell(r, 3).value or "").strip()
    ad = kanonik.get(no)
    if ad is None:
        eslesmedi.append({"no": no, "adKaynak": ad_kaynak})
        continue
    pot, rez = {}, {}
    for y, c in yil_kol.items():
        b = r3(ws.cell(r, c).value)
        i = r3(ws.cell(r, c + 1).value)
        if b is not None:
            pot[str(y)] = b
        if i is not None:
            rez[str(y)] = i
    havzalar[no] = {"ad": ad, "adKaynak": ad_kaynak, "potansiyel": pot, "rezerv": rez}

# --- 1.5: ulusal rezerv serisi (yıl=col2, değer=col3) ---
w5 = openpyxl.load_workbook(F15, data_only=True).active
ulusalRezerv = {}
for r in range(1, w5.max_row + 1):
    y = w5.cell(r, 2).value
    v = r3(w5.cell(r, 3).value)
    if isinstance(y, int) and 1990 <= y <= 2030 and v is not None:
        ulusalRezerv[str(y)] = v

# --- 1.6: ulusal sektör tahsisi (yıl=col2; sulama=col5, içme/sanayi=col6, toplam=col7) ---
w6 = openpyxl.load_workbook(F16, data_only=True).active
tahsisToplamSeri = {}
tahsisSonYil = None
for r in range(1, w6.max_row + 1):
    y = w6.cell(r, 2).value
    top = r3(w6.cell(r, 7).value)
    if isinstance(y, int) and 1990 <= y <= 2030 and top is not None:
        tahsisToplamSeri[str(y)] = top
        tahsisSonYil = {
            "yil": y,
            "toplam": top,
            "sulama": r3(w6.cell(r, 5).value),
            "icmeSanayi": r3(w6.cell(r, 6).value),
        }

cikti = {
    "aciklama": "Türkiye 25 havzası için resmî DSİ yeraltısuyu (YAS) verisi — "
    "havza potansiyel bloğunun TEK kaynağı. Boş/atlanan alan = xlsx'te yok; "
    "uydurma/enterpolasyon yapılmaz. Eşleşme resmî havza no (1-25) üzerinden.",
    "birim": "hm³/yıl (havza) · km³/yıl (ulusal)",
    "yillar": yillar,
    "kaynak": {
        "havzaPotansiyelRezerv": {
            "ad": "DSİ Resmî Su Kaynakları İstatistikleri, Tablo 1.3 "
            "(Havzalara Göre Yıllık Yeraltısuyu Potansiyeli, 2013-2024)",
            "dosya": "1.3.havzalara_gore_yillik_yeraltisuyu_potansiyeli_20132024.xlsx",
            "url": URL, "indirmeTarihi": TARIH,
        },
        "ulusalRezerv": {
            "ad": "DSİ Resmî Su Kaynakları İstatistikleri, Tablo 1.5 "
            "(Yeraltısuyu İşletme Rezervi, 1995-2024)",
            "dosya": "1.5.yeraltisuyu_isletme_rezervi_19952024.xlsx",
            "url": URL, "indirmeTarihi": TARIH,
        },
        "ulusalTahsis": {
            "ad": "DSİ Resmî Su Kaynakları İstatistikleri, Tablo 1.6 "
            "(Sektör Bazında Yeraltısuyundan Yapılan Tahsis Miktarı, 1995-2024)",
            "dosya": "1.6.sektor_bazinda_yeraltisuyundan_yapilan_tahsis_miktari_19952024.xlsx",
            "url": URL, "indirmeTarihi": TARIH,
        },
    },
    "ulusal": {
        "rezervBirim": "km³/yıl",
        "rezervSeri": ulusalRezerv,
        "tahsisBirim": "km³/yıl",
        "tahsisToplamSeri": tahsisToplamSeri,
        "tahsisSonYil": tahsisSonYil,
    },
    "havzalar": havzalar,
    "eslesmedi": eslesmedi,
}

json.dump(cikti, open(CIKTI, "w"), ensure_ascii=False, indent=2)
print(f"yazıldı: {CIKTI}")
print(f"havza eşleşen: {len(havzalar)}/25 · eşleşmedi: {len(eslesmedi)}")
print(f"yıllar: {yillar[0]}-{yillar[-1]} ({len(yillar)})")
print(f"ulusal rezerv yıl aralığı: {min(ulusalRezerv)}-{max(ulusalRezerv)}")
print(f"tahsis son yıl: {tahsisSonYil}")
# Sakarya (12) örnek
s = havzalar.get("12")
if s:
    print(f"Sakarya potansiyel 2013/2024: {s['potansiyel'].get('2013')}/{s['potansiyel'].get('2024')} "
          f"· rezerv 2024: {s['rezerv'].get('2024')}")
